import json

from extract import commit, scan


def make_dirs(tmp_path):
    raw = tmp_path / "raw"
    raw.mkdir()
    return raw, tmp_path / "extracted", tmp_path / "manifest.json"


def test_new_markdown_file_is_extracted_and_recorded_in_manifest_on_commit(tmp_path):
    raw, extracted, manifest = make_dirs(tmp_path)
    (raw / "weekly").mkdir()
    (raw / "weekly" / "w1.md").write_text("本週完成 WGC 整合\n", encoding="utf-8")

    result = scan(raw, extracted, manifest)

    assert result.new == ["weekly/w1.md"]
    assert result.changed == []
    assert (extracted / "weekly" / "w1.md.txt").read_text(encoding="utf-8") == "本週完成 WGC 整合\n"
    assert not manifest.exists()
    commit(raw, manifest, result.new)
    entry = json.loads(manifest.read_text(encoding="utf-8"))["weekly/w1.md"]
    assert len(entry["sha256"]) == 64
    assert entry["scanned_at"]


def test_rerun_on_unchanged_files_processes_nothing(tmp_path):
    raw, extracted, manifest = make_dirs(tmp_path)
    (raw / "a.md").write_text("alpha", encoding="utf-8")
    commit(raw, manifest, scan(raw, extracted, manifest).new)
    first_scanned_at = json.loads(manifest.read_text(encoding="utf-8"))["a.md"]["scanned_at"]
    (extracted / "a.md.txt").unlink()

    result = scan(raw, extracted, manifest)

    assert result.new == []
    assert result.changed == []
    assert not (extracted / "a.md.txt").exists()
    assert json.loads(manifest.read_text(encoding="utf-8"))["a.md"]["scanned_at"] == first_scanned_at


def test_modified_file_is_reported_as_changed_and_reextracted(tmp_path):
    raw, extracted, manifest = make_dirs(tmp_path)
    (raw / "a.md").write_text("alpha", encoding="utf-8")
    (raw / "b.md").write_text("beta", encoding="utf-8")
    commit(raw, manifest, scan(raw, extracted, manifest).new)

    (raw / "a.md").write_text("alpha v2", encoding="utf-8")
    (raw / "c.md").write_text("gamma", encoding="utf-8")
    result = scan(raw, extracted, manifest)

    assert result.changed == ["a.md"]
    assert result.new == ["c.md"]
    assert (extracted / "a.md.txt").read_text(encoding="utf-8") == "alpha v2"


def test_uncommitted_files_are_reported_again_on_next_scan(tmp_path):
    raw, extracted, manifest = make_dirs(tmp_path)
    (raw / "a.md").write_text("alpha", encoding="utf-8")
    (raw / "b.md").write_text("beta", encoding="utf-8")

    first = scan(raw, extracted, manifest)
    second = scan(raw, extracted, manifest)

    assert first.new == second.new == ["a.md", "b.md"]


def test_commit_marks_only_given_files_and_modification_reopens_them(tmp_path):
    raw, extracted, manifest = make_dirs(tmp_path)
    (raw / "a.md").write_text("alpha", encoding="utf-8")
    (raw / "b.md").write_text("beta", encoding="utf-8")
    scan(raw, extracted, manifest)

    commit(raw, manifest, ["a.md"])
    assert scan(raw, extracted, manifest).new == ["b.md"]

    commit(raw, manifest, ["b.md"])
    result = scan(raw, extracted, manifest)
    assert result.new == [] and result.changed == []

    (raw / "a.md").write_text("alpha v2", encoding="utf-8")
    assert scan(raw, extracted, manifest).changed == ["a.md"]


def extracted_text(tmp_path, name):
    return (tmp_path / "extracted" / (name + ".txt")).read_text(encoding="utf-8")


def test_docx_paragraphs_and_table_cells_are_extracted(tmp_path):
    from docx import Document

    raw, extracted, manifest = make_dirs(tmp_path)
    doc = Document()
    doc.add_paragraph("年度考績 2024")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "KPI"
    table.cell(0, 1).text = "crash rate -30%"
    doc.save(raw / "review.docx")

    scan(raw, extracted, manifest)

    text = extracted_text(tmp_path, "review.docx")
    assert "年度考績 2024" in text
    assert "KPI" in text and "crash rate -30%" in text


def test_pptx_slide_text_and_notes_are_extracted(tmp_path):
    from pptx import Presentation

    raw, extracted, manifest = make_dirs(tmp_path)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "ISE 2023 展示"
    slide.notes_slide.notes_text_frame.text = "Demo speaker note"
    prs.save(raw / "demo.pptx")

    scan(raw, extracted, manifest)

    text = extracted_text(tmp_path, "demo.pptx")
    assert "ISE 2023 展示" in text
    assert "Demo speaker note" in text


def test_xlsx_cells_are_extracted_per_sheet(tmp_path):
    from openpyxl import Workbook

    raw, extracted, manifest = make_dirs(tmp_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"
    ws.append(["metric", "value"])
    ws.append(["crash rate", 0.7])
    wb.create_sheet("Notes").append(["季度回顧"])
    wb.save(raw / "kpi.xlsx")

    scan(raw, extracted, manifest)

    text = extracted_text(tmp_path, "kpi.xlsx")
    assert "KPI" in text and "crash rate\t0.7" in text
    assert "Notes" in text and "季度回顧" in text


def write_minimal_pdf(path, text):
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(out)


def test_pdf_page_text_is_extracted(tmp_path):
    raw, extracted, manifest = make_dirs(tmp_path)
    write_minimal_pdf(raw / "old-resume.pdf", "Senior Software Engineer 2019-2024")

    scan(raw, extracted, manifest)

    text = extracted_text(tmp_path, "old-resume.pdf")
    assert "Senior Software Engineer 2019-2024" in text
    assert "%PDF" not in text and "endobj" not in text


def test_attachments_folders_and_unsupported_formats_are_skipped(tmp_path):
    raw, extracted, manifest = make_dirs(tmp_path)
    (raw / "wiki" / "attachments").mkdir(parents=True)
    (raw / "wiki" / "page.md").write_text("page", encoding="utf-8")
    (raw / "wiki" / "attachments" / "notes.txt").write_text("noise", encoding="utf-8")
    (raw / "wiki" / "attachments" / "shot.png").write_bytes(b"\x89PNG\x00\xff")
    (raw / "demo.mov").write_bytes(b"\x00\x01\xff\xfe")
    (raw / "tool.EXE").write_bytes(b"MZ\x90\x00\xff")

    result = scan(raw, extracted, manifest)

    assert result.new == ["wiki/page.md"]
    assert result.skipped == ["demo.mov", "tool.EXE"]
    assert not manifest.exists()
    assert not (extracted / "wiki" / "attachments").exists()

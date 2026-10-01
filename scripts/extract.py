"""Extract Raw Records into plain text and track them in a hash manifest."""
import hashlib
import json
import sys
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path


@dataclass
class ScanResult:
    new: list[str] = field(default_factory=list)
    changed: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)  # unsupported formats


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    lines = []
    for slide in Presentation(path).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
        if slide.has_notes_slide:
            lines.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(lines)


def _xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                lines.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(lines)


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


TEXT_SUFFIXES = {".md", ".txt"}
IGNORED_DIRS = {"attachments"}
_EXTRACTORS = {
    ".docx": _docx_text,
    ".pptx": _pptx_text,
    ".xlsx": _xlsx_text,
    ".pdf": _pdf_text,
}


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8")
    return _EXTRACTORS[suffix](path)


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in TEXT_SUFFIXES | _EXTRACTORS.keys()


def _load_manifest(manifest_path: Path) -> dict:
    if manifest_path.exists():
        return json.loads(manifest_path.read_text(encoding="utf-8"))
    return {}


def scan(raw_dir: Path, extracted_dir: Path, manifest_path: Path) -> ScanResult:
    """Extract new/changed files. The manifest is read-only here: files only count
    as scanned once `commit` is called after their content was merged."""
    manifest = _load_manifest(manifest_path)
    result = ScanResult()
    for path in sorted(raw_dir.rglob("*")):
        if not path.is_file():
            continue
        rel = path.relative_to(raw_dir).as_posix()
        if IGNORED_DIRS & set(path.relative_to(raw_dir).parts[:-1]):
            continue
        if not is_supported(path):
            result.skipped.append(rel)
            continue
        digest = _sha256(path)
        known = manifest.get(rel)
        if known and known["sha256"] == digest:
            continue
        out = extracted_dir / (rel + ".txt")
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(extract_text(path), encoding="utf-8")
        (result.changed if known else result.new).append(rel)
    return result


def commit(raw_dir: Path, manifest_path: Path, rels: list[str]) -> None:
    """Record files as scanned, once their content is merged and validated."""
    manifest = _load_manifest(manifest_path)
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    for rel in rels:
        path = raw_dir / rel
        if not path.is_file():
            raise FileNotFoundError(f"not a Raw Record: {rel}")
        manifest[rel] = {"sha256": _sha256(path), "scanned_at": now}
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw", type=Path, default=Path("data/raw"))
    parser.add_argument("--out", type=Path, default=Path("data/extracted"))
    parser.add_argument("--manifest", type=Path, default=Path("data/manifest.json"))
    parser.add_argument(
        "--commit",
        nargs="+",
        metavar="PATH",
        help="mark these files (relative to --raw, as printed by extract) as scanned",
    )
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")  # wiki paths contain non-ASCII characters

    if args.commit:
        commit(args.raw, args.manifest, args.commit)
        print(f"committed {len(args.commit)} file(s)")
        return

    result = scan(args.raw, args.out, args.manifest)
    for label, files in (("new", result.new), ("changed", result.changed)):
        for rel in files:
            print(f"{label}\t{rel}")
    print(
        f"{len(result.new)} new, {len(result.changed)} changed, "
        f"{len(result.skipped)} skipped (unsupported format)"
    )


if __name__ == "__main__":
    main()

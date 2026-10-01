"""Generate templates/cv-clean.docx, the plain black-and-white CV Template (docxtpl tags)."""
import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Calibri"
BODY_PT = 10.5


def _para(doc, text="", size=BODY_PT, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
          before=0, after=0, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
    return p


def _right_tab(p, width_cm):
    p.paragraph_format.tab_stops.add_tab_stop(Cm(width_cm), WD_TAB_ALIGNMENT.RIGHT)


def _bottom_border(p):
    pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "000000")):
        bottom.set(qn(f"w:{k}"), v)
    borders.append(bottom)
    pr.insert_element_before(
        borders, "w:shd", "w:tabs", "w:spacing", "w:ind", "w:contextualSpacing", "w:jc", "w:rPr")


def _heading(doc, text):
    p = _para(doc, text, size=11.5, bold=True, before=10, after=3)
    p.paragraph_format.keep_with_next = True
    _bottom_border(p)


def _tag(doc, tag):
    """A docxtpl paragraph-level control tag ({%p ... %}) that renders to nothing."""
    return _para(doc, tag)


def build(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    width = 17.0  # text width in cm

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(BODY_PT)

    # Header: name, headline, contact line
    _para(doc, "{{ basics.name }}", size=20, bold=True, after=0)
    _para(doc, "{{ basics.headline }}", size=11.5, after=2)
    _para(doc, "{{ basics.contact }}", after=4)

    # Summary
    _tag(doc, "{%p if summary %}")
    _heading(doc, "SUMMARY")
    _para(doc, "{{ summary }}", after=2)
    _tag(doc, "{%p endif %}")

    # Skills, one line per category
    _tag(doc, "{%p if skills %}")
    _heading(doc, "SKILLS")
    _tag(doc, "{%p for g in skills %}")
    p = _para(doc, after=1)
    p.paragraph_format.left_indent = Cm(3.2)
    p.paragraph_format.first_line_indent = Cm(-3.2)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(3.2))
    r = p.add_run("{{ g.category }}")
    r.bold = True
    p.add_run("\t{{ g.names }}")
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p endif %}")

    # Experience
    _heading(doc, "WORK EXPERIENCE")
    _tag(doc, "{%p for job in experience %}")
    p = _para(doc, before=5, after=0)
    _right_tab(p, width)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("{{ job.company }}")
    r.bold = True
    p.add_run("{{ job.location_suffix }}")
    r = p.add_run("\t{{ job.period }}")
    r.bold = True
    p = _para(doc, "{{ job.title }}", after=2)
    p.runs[0].italic = True
    p.paragraph_format.keep_with_next = True
    _tag(doc, "{%p for b in job.bullets %}")
    p = _para(doc, "{{ b }}", after=1, style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p endfor %}")

    # Education
    _tag(doc, "{%p if education %}")
    _heading(doc, "EDUCATION")
    _tag(doc, "{%p for e in education %}")
    p = _para(doc, before=2)
    _right_tab(p, width)
    r = p.add_run("{{ e.school }}")
    r.bold = True
    p.add_run("\t{{ e.period }}")
    _para(doc, "{{ e.degree }}")
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p endif %}")

    # Certifications
    _tag(doc, "{%p if certifications %}")
    _heading(doc, "CERTIFICATIONS")
    _tag(doc, "{%p for c in certifications %}")
    p = _para(doc, "{{ c.line }}", after=1)
    _tag(doc, "{%p endfor %}")
    _tag(doc, "{%p endif %}")

    # Languages
    _tag(doc, "{%p if languages %}")
    _heading(doc, "LANGUAGES")
    _para(doc, "{{ languages }}")
    _tag(doc, "{%p endif %}")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="templates/cv-clean.docx")
    build(Path(ap.parse_args().out))
    print("wrote", ap.parse_args().out)
